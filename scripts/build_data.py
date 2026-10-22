"""Build data/housing-data.json for the NZ housing affordability calculator.

Reads inbox/housing-inputs.xlsx (updated by hand once a year):
  "House prices"      REINZ median house price for August, one column per year from 2005
  "Household income"  Stats NZ median WEEKLY household income (year ended June), one column per year from 2003
Interest rates (1-year fixed, August) come from the nz-mortgage-rates project's published files, so
they never need typing. If that download fails, the rates already in data/housing-data.json are reused.

The script stops with an error, and publishes nothing, if anything looks wrong.
"""
import csv, datetime, io, json, os, sys, urllib.request
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
INPUTS = os.path.join(ROOT, 'inbox', 'housing-inputs.xlsx')
OUT = os.path.join(ROOT, 'data', 'housing-data.json')
RATES_STANDARD = 'https://carrnick308.github.io/nz-mortgage-rates/nz-mortgage-rates.csv'
RATES_SPECIAL = 'https://carrnick308.github.io/nz-mortgage-rates/nz-mortgage-rates-b21.csv'
REGIONS = ['New Zealand', 'Northland', 'Auckland', 'Waikato', 'Bay of Plenty', 'Taranaki',
           'Manawatu', 'Wellington', 'Canterbury', 'Otago', 'Southland']
FIRST_YEAR = 2005        # first year shown on the page
SPECIAL_FROM = 2017      # RBNZ special rates (B21) start in 2017


def fail(msg):
    print('ERROR: ' + msg)
    sys.exit(1)


def warn(msg):
    print('WARNING: ' + msg)


def read_sheet(wb, name, lo, hi, what):
    if name not in wb.sheetnames:
        fail('The sheet "%s" is missing from housing-inputs.xlsx' % name)
    ws = wb[name]
    rows = [r for r in ws.iter_rows(values_only=True)]
    head = next((r for r in rows if r and r[0] == 'Region'), None)
    if head is None:
        fail('Sheet "%s": no row starting with "Region" (the year headings)' % name)
    years = []
    for v in head[1:]:
        if v is None or v == '':
            break
        try:
            years.append(int(v))
        except (TypeError, ValueError):
            fail('Sheet "%s": "%s" in the year headings is not a year' % (name, v))
    if years != list(range(years[0], years[0] + len(years))):
        fail('Sheet "%s": the years must run one after another with no gaps (found %s)' % (name, years))
    data = {}
    for r in rows:
        if r and isinstance(r[0], str) and r[0].strip() in REGIONS:
            label = r[0].strip()
            vals = list(r[1:1 + len(years)])
            for y, v in zip(years, vals):
                if v is None or v == '':
                    fail('Sheet "%s": %s has no figure for %d' % (name, label, y))
                if isinstance(v, str):
                    try:
                        v = float(v.replace(',', '').replace('$', ''))
                    except ValueError:
                        fail('Sheet "%s": %s %d is not a number ("%s")' % (name, label, y, v))
                if not (lo <= v <= hi):
                    fail('Sheet "%s": %s %d is %s, which looks wrong for %s' % (name, label, y, v, what))
            data[label] = dict(zip(years, [float(str(v).replace(',', '').replace('$', '')) for v in vals]))
    missing = [r for r in REGIONS if r not in data]
    if missing:
        fail('Sheet "%s": these rows are missing or misspelt: %s' % (name, ', '.join(missing)))
    return years, data


def august(url):
    with urllib.request.urlopen(url, timeout=30) as resp:
        text = resp.read().decode('utf-8')
    out = {}
    for row in csv.DictReader(io.StringIO(text)):
        if row.get('Month', '').endswith('-08') and row.get('1 year'):
            out[int(row['Month'][:4])] = float(row['1 year'])
    if not out:
        raise ValueError('no August rows in ' + url)
    return out


if not os.path.exists(INPUTS):
    fail('inbox/housing-inputs.xlsx not found')
wb = openpyxl.load_workbook(INPUTS, data_only=True)
p_years, prices = read_sheet(wb, 'House prices', 50000, 5000000, 'a median house price')
i_years, incomes = read_sheet(wb, 'Household income', 300, 15000, 'a WEEKLY household income')

if p_years[0] > FIRST_YEAR:
    fail('House prices must start in %d or earlier' % FIRST_YEAR)
if i_years[0] > FIRST_YEAR - 2:
    fail('Household income must start in %d or earlier (the 3-year average needs two years before %d)' % (FIRST_YEAR - 2, FIRST_YEAR))
last = min(p_years[-1], i_years[-1])
if p_years[-1] != i_years[-1]:
    warn('House prices run to %d but incomes to %d. The page will stop at %d until both are in.' % (p_years[-1], i_years[-1], last))
years = list(range(FIRST_YEAR, last + 1))

# sanity: big jumps are allowed (small regions move a lot) but flagged in the log
for r in REGIONS:
    for y in years[1:]:
        ch = prices[r][y] / prices[r][y - 1] - 1
        if abs(ch) > 0.4:
            warn('%s house price %d moved %.0f%% from %d. Check it was typed correctly.' % (r, y, ch * 100, y - 1))

# interest rates
old = json.load(open(OUT)) if os.path.exists(OUT) else None
try:
    std, spc = august(RATES_STANDARD), august(RATES_SPECIAL)
    print('Downloaded interest rates from nz-mortgage-rates')
except Exception as e:  # fall back to the rates already published
    warn('Could not download interest rates (%s). Reusing the rates already published.' % e)
    if not old:
        fail('No published rates to fall back on')
    std = {y: v for y, v in zip(old['years'], old['rateStandard']) if v is not None}
    spc = {y: v for y, v in zip(old['years'], old['rateSpecial']) if v is not None}
for y in years:
    if y not in std:
        fail('No standard 1-year rate for August %d yet. Wait until the Reserve Bank publishes August, then upload again.' % y)
    if y >= SPECIAL_FROM and y not in spc:
        fail('No special 1-year rate for August %d yet. Wait until the Reserve Bank publishes August, then upload again.' % y)

key = lambda r: 'National' if r == 'New Zealand' else r
data = {
    'asOf': datetime.datetime.now(datetime.timezone.utc).strftime('%Y-%m-%d'),
    'years': years,
    'regions': [key(r) for r in REGIONS],
    'price': {key(r): [round(prices[r][y]) for y in years] for r in REGIONS},
    'incomeFrom': FIRST_YEAR - 2,
    'incomeWeekly': {key(r): [round(incomes[r][y]) for y in range(FIRST_YEAR - 2, last + 1)] for r in REGIONS},
    'rateStandard': [std[y] for y in years],
    'rateSpecial': [spc.get(y) if y >= SPECIAL_FROM else None for y in years],
}
if old and len(data['years']) < len(old['years']):
    fail('The new data has fewer years (%d) than what is published (%d). Nothing was changed.' % (len(data['years']), len(old['years'])))
os.makedirs(os.path.dirname(OUT), exist_ok=True)
with open(OUT, 'w') as f:
    json.dump(data, f, separators=(',', ':'))
print('Built housing-data.json: %d to %d, %d regions' % (years[0], years[-1], len(REGIONS)))
